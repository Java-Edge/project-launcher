#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Dump the shapes of one or all slides in a PPTX as JSON.

Use this before writing a fill spec: it shows which slide index has the
title+table style you want to clone, what the current title/header/sample
row text look like, and how many rows/columns the template table has.

Usage:
    python3 inspect_pptx_slide.py deck.pptx
    python3 inspect_pptx_slide.py deck.pptx --slide-index 0
"""
from __future__ import annotations

import argparse
import json

from pptx import Presentation


def describe(path: str, slide_index: int | None):
    prs = Presentation(path)
    indices = [slide_index] if slide_index is not None else range(len(prs.slides))
    results = []
    for i in indices:
        slide = prs.slides[i]
        info = {"index": i, "layout": slide.slide_layout.name, "shapes": []}
        for shp in slide.shapes:
            entry = {"name": shp.name, "shape_type": str(shp.shape_type)}
            if shp.has_table:
                tbl = shp.table
                n_cols = len(tbl.columns)
                entry["table_rows"] = len(tbl.rows)
                entry["table_cols"] = n_cols
                entry["header_row"] = [tbl.cell(0, c).text for c in range(n_cols)]
                if len(tbl.rows) > 1:
                    entry["sample_body_row"] = [tbl.cell(1, c).text for c in range(n_cols)]
            elif shp.has_text_frame:
                entry["text"] = shp.text_frame.text
                first_para_runs = shp.text_frame.paragraphs[0].runs if shp.text_frame.paragraphs else []
                entry["first_paragraph_run_count"] = len(first_para_runs)
                entry["first_paragraph_run_texts"] = [r.text for r in first_para_runs]
            info["shapes"].append(entry)
        results.append(info)
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx", help="Path to the PPTX file")
    ap.add_argument("--slide-index", type=int, default=None, help="Inspect only this 0-based slide index")
    args = ap.parse_args()
    describe(args.pptx, args.slide_index)
