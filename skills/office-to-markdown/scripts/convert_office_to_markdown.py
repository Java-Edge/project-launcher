#!/usr/bin/env python3
"""Convert DOCX/PPTX/XLSX/XLS to Markdown with local image assets.

This script intentionally wraps MarkItDown instead of calling it raw. MarkItDown
can emit images as data URIs; this script saves those images to disk and rewrites
the Markdown links so the result is portable and readable.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any


SUPPORTED_EXTENSIONS = {".docx", ".pptx", ".xlsx", ".xls"}
ZIP_IMAGE_PREFIX = {
    ".docx": "word/media/",
    ".pptx": "ppt/media/",
    ".xlsx": "xl/media/",
}
IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".bmp",
    ".tif",
    ".tiff",
    ".svg",
    ".webp",
    ".emf",
    ".wmf",
}
MIME_EXTENSION = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/jpg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
    "image/svg+xml": ".svg",
    "image/webp": ".webp",
    "image/x-emf": ".emf",
    "image/x-wmf": ".wmf",
}
DATA_IMAGE_RE = re.compile(
    r"!\[([^\]]*)\]\(data:(image/[A-Za-z0-9.+-]+);base64,([A-Za-z0-9+/=\s]+)\)"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Office documents to Markdown with extracted image assets."
    )
    parser.add_argument("input", help="Source .docx, .pptx, .xlsx, or .xls file")
    parser.add_argument("--out", help="Output Markdown path")
    parser.add_argument("--asset-dir", help="Directory for extracted image assets")
    parser.add_argument("--report-out", help="JSON quality report path")
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Fail if unsupported features or unplaced images are detected",
    )
    parser.add_argument(
        "--no-append-unplaced",
        action="store_true",
        help="Extract package images but do not append images that have no inline position",
    )
    parser.add_argument(
        "--pptx-engine",
        choices=["structured", "markitdown"],
        default="structured",
        help="Use a structure-aware python-pptx converter for PPTX, or raw MarkItDown",
    )
    parser.add_argument(
        "--pptx-list-mode",
        choices=["auto", "aggressive", "off"],
        default="auto",
        help="How strongly PPTX paragraph groups should become Markdown lists",
    )
    return parser.parse_args()


def sanitize_filename(value: str, fallback: str) -> str:
    value = value.strip().replace("\x00", "")
    value = re.sub(r"[\\/:*?\"<>|]+", "-", value)
    value = re.sub(r"\s+", "-", value)
    value = re.sub(r"-+", "-", value).strip(".-")
    return value[:80] or fallback


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    parent = path.parent
    for index in range(2, 10_000):
        candidate = parent / f"{stem}-{index}{suffix}"
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"unable to allocate unique path for {path}")


def rel_link(target: Path, markdown_path: Path) -> str:
    rel = os.path.relpath(target, start=markdown_path.parent)
    rel = Path(rel).as_posix()
    if re.search(r"[\s()<>]", rel):
        return f"<{rel}>"
    return rel


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def markitdown_convert(input_path: Path) -> tuple[str, dict[str, Any]]:
    cli = shutil.which("markitdown")
    if cli:
        proc = subprocess.run(
            [cli, "--keep-data-uris", str(input_path)],
            check=False,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        if proc.returncode == 0:
            return proc.stdout, {
                "method": "cli",
                "path": cli,
                "stderr": proc.stderr.strip(),
            }
        cli_error = proc.stderr.strip() or proc.stdout.strip()
    else:
        cli_error = "markitdown CLI not found"

    try:
        from markitdown import MarkItDown

        result = MarkItDown().convert(str(input_path), keep_data_uris=True)
        return result.markdown, {
            "method": "python-api",
            "path": "markitdown.MarkItDown",
            "cli_error": cli_error,
        }
    except Exception as exc:  # pragma: no cover - depends on local install
        raise RuntimeError(f"MarkItDown conversion failed: {cli_error}; {exc}") from exc


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]

    def cell(value: str) -> str:
        value = re.sub(r"\s*\n\s*", "<br>", value.strip())
        return value.replace("|", "\\|")

    lines = ["| " + " | ".join(cell(v) for v in normalized[0]) + " |"]
    lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
    for row in normalized[1:]:
        lines.append("| " + " | ".join(cell(v) for v in row) + " |")
    return "\n".join(lines)


def placeholder_type_name(shape: Any) -> str:
    try:
        if shape.is_placeholder:
            return str(shape.placeholder_format.type).split(" (", 1)[0].upper()
    except Exception:
        pass
    return ""


def shape_sort_key(shape: Any) -> tuple[int, int]:
    top = getattr(shape, "top", None)
    left = getattr(shape, "left", None)
    return (int(top or 0), int(left or 0))


def iter_pptx_shapes(shapes: Any) -> list[Any]:
    ordered: list[Any] = []
    for shape in sorted(shapes, key=shape_sort_key):
        if hasattr(shape, "shapes") and "GROUP" in str(getattr(shape, "shape_type", "")):
            ordered.extend(iter_pptx_shapes(shape.shapes))
        else:
            ordered.append(shape)
    return ordered


def same_pptx_shape(left: Any, right: Any | None) -> bool:
    if right is None:
        return False
    if left is right:
        return True
    return getattr(left, "_element", None) is getattr(right, "_element", None)


def clean_text(value: str) -> str:
    value = value.replace("\u00a0", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\s+\n", "\n", value)
    return value.strip()


def clean_bullet_prefix(value: str) -> str:
    value = clean_text(value)
    return re.sub(r"^[\u2022\u25cf\u25cb\u25a0\u25aa\u2219\-–—]\s*", "", value).strip()


def paragraph_font_size(paragraph: Any) -> float | None:
    sizes: list[float] = []
    try:
        if paragraph.font.size is not None:
            sizes.append(float(paragraph.font.size.pt))
    except Exception:
        pass
    for run in getattr(paragraph, "runs", []):
        try:
            if run.font.size is not None:
                sizes.append(float(run.font.size.pt))
        except Exception:
            pass
    return max(sizes) if sizes else None


def paragraph_is_bold(paragraph: Any) -> bool:
    try:
        if paragraph.font.bold:
            return True
    except Exception:
        pass
    for run in getattr(paragraph, "runs", []):
        try:
            if run.font.bold:
                return True
        except Exception:
            pass
    return False


def paragraph_has_explicit_bullet(paragraph: Any) -> bool:
    xml = getattr(getattr(paragraph, "_p", None), "xml", "")
    if "<a:buNone" in xml:
        return False
    return any(token in xml for token in ("<a:buChar", "<a:buAutoNum", "<a:buBlip"))


def nonempty_paragraphs(shape: Any) -> list[Any]:
    if not getattr(shape, "has_text_frame", False):
        return []
    return [p for p in shape.text_frame.paragraphs if clean_text(getattr(p, "text", ""))]


def paragraph_record(paragraph: Any) -> dict[str, Any]:
    return {
        "text": clean_text(paragraph.text),
        "level": int(getattr(paragraph, "level", 0) or 0),
        "font_size": paragraph_font_size(paragraph),
        "bold": paragraph_is_bold(paragraph),
        "explicit_bullet": paragraph_has_explicit_bullet(paragraph),
    }


def heading_candidate_score(shape: Any, slide_height: int) -> tuple[float, int]:
    paragraphs = nonempty_paragraphs(shape)
    if not paragraphs:
        return (0.0, 0)
    font_size = max((paragraph_font_size(p) or 0.0) for p in paragraphs)
    top = int(getattr(shape, "top", 0) or 0)
    placeholder = placeholder_type_name(shape)
    score = font_size
    if placeholder in {"TITLE", "CENTER_TITLE"}:
        score += 100
    if top <= slide_height * 0.35:
        score += 12
    return (score, -top)


def choose_pptx_title_shape(slide: Any, slide_height: int) -> Any | None:
    try:
        if slide.shapes.title is not None and clean_text(slide.shapes.title.text):
            return slide.shapes.title
    except Exception:
        pass
    candidates = [
        shape
        for shape in iter_pptx_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False) and nonempty_paragraphs(shape)
    ]
    if not candidates:
        return None
    best = max(candidates, key=lambda shape: heading_candidate_score(shape, slide_height))
    score, _ = heading_candidate_score(best, slide_height)
    if score >= 24:
        return best
    return None


def is_body_placeholder(shape: Any) -> bool:
    placeholder = placeholder_type_name(shape)
    return any(
        key in placeholder
        for key in ("BODY", "OBJECT", "CONTENT", "TEXT", "VERTICAL_BODY")
    )


def first_paragraph_is_section_heading(records: list[dict[str, Any]], body_like: bool) -> bool:
    if len(records) < 2:
        return False
    first = records[0]
    rest_sizes = [r["font_size"] for r in records[1:] if r["font_size"] is not None]
    rest_max = max(rest_sizes) if rest_sizes else None
    if first["explicit_bullet"] or first["level"] > 0:
        return False
    if first["bold"] and not body_like:
        return True
    if first["font_size"] and rest_max and first["font_size"] >= rest_max + 2:
        return True
    return False


def should_render_as_list(
    record: dict[str, Any],
    records: list[dict[str, Any]],
    body_like: bool,
    list_mode: str,
) -> bool:
    if list_mode == "off":
        return False
    if record["explicit_bullet"] or record["level"] > 0:
        return True
    if list_mode == "aggressive" and len(records) > 1:
        return True
    if body_like and len(records) > 1:
        return True
    if len(records) >= 2 and all(len(item["text"]) <= 140 for item in records):
        return True
    return False


def pptx_text_shape_to_markdown(
    shape: Any,
    *,
    is_title_shape: bool,
    list_mode: str,
) -> tuple[list[str], dict[str, int]]:
    records = [paragraph_record(p) for p in nonempty_paragraphs(shape)]
    if not records:
        return [], {"headings": 0, "list_items": 0, "body_paragraphs": 0}

    placeholder = placeholder_type_name(shape)
    body_like = is_body_placeholder(shape)
    lines: list[str] = []
    stats = {"headings": 0, "list_items": 0, "body_paragraphs": 0}

    if is_title_shape:
        return [], stats

    if placeholder == "SUBTITLE":
        text = " ".join(record["text"] for record in records)
        return [f"### {text}"], {"headings": 1, "list_items": 0, "body_paragraphs": 0}

    start = 0
    section_heading_emitted = False
    if first_paragraph_is_section_heading(records, body_like):
        lines.append(f"### {records[0]['text']}")
        lines.append("")
        stats["headings"] += 1
        start = 1
        section_heading_emitted = True

    list_records = records[start:] or records
    for record in records[start:]:
        text = clean_bullet_prefix(record["text"])
        if not text:
            continue
        if should_render_as_list(
            record,
            list_records,
            body_like or section_heading_emitted,
            list_mode,
        ):
            indent = "  " * max(0, min(int(record["level"]), 6))
            lines.append(f"{indent}- {text}")
            stats["list_items"] += 1
        else:
            lines.append(text)
            lines.append("")
            stats["body_paragraphs"] += 1

    while lines and lines[-1] == "":
        lines.pop()
    return lines, stats


def save_pptx_image(
    shape: Any,
    asset_dir: Path,
    markdown_path: Path,
    image_hashes: set[str],
    digest_to_path: dict[str, Path],
    index: int,
) -> tuple[str, dict[str, Any]]:
    image = shape.image
    raw = image.blob
    digest = sha256_bytes(raw)
    image_hashes.add(digest)

    if digest in digest_to_path:
        out_path = digest_to_path[digest]
    else:
        ext = "." + (getattr(image, "ext", "") or "").lower().lstrip(".")
        if ext == ".":
            ext = MIME_EXTENSION.get(getattr(image, "content_type", "").lower(), ".bin")
        source_name = Path(getattr(image, "filename", "") or "").stem
        if not source_name:
            source_name = getattr(shape, "name", "") or f"image-{index:03d}"
        stem = sanitize_filename(source_name, f"image-{index:03d}")
        out_path = unique_path(asset_dir / f"{stem}{ext}")
        out_path.write_bytes(raw)
        digest_to_path[digest] = out_path

    try:
        alt = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
    except Exception:
        alt = ""
    alt = clean_text(alt or getattr(shape, "name", "") or f"image {index:03d}")
    alt = re.sub(r"[\[\]\r\n]+", " ", alt).strip()
    link = rel_link(out_path, markdown_path)
    return (
        f"![{alt}]({link})",
        {
            "alt": alt,
            "mime": getattr(image, "content_type", ""),
            "path": str(out_path),
            "markdown_link": link,
            "bytes": len(raw),
            "sha256": digest,
            "slide_position": {
                "top": int(getattr(shape, "top", 0) or 0),
                "left": int(getattr(shape, "left", 0) or 0),
            },
        },
    )


def pptx_table_to_markdown(shape: Any) -> str:
    rows: list[list[str]] = []
    for row in shape.table.rows:
        rows.append([clean_text(cell.text) for cell in row.cells])
    return markdown_table(rows)


def pptx_chart_title(shape: Any) -> str:
    try:
        chart = shape.chart
        if chart.has_title:
            return clean_text(chart.chart_title.text_frame.text)
    except Exception:
        pass
    return clean_text(getattr(shape, "name", "") or "Chart")


def pptx_structured_convert(
    input_path: Path,
    asset_dir: Path,
    markdown_path: Path,
    list_mode: str,
) -> tuple[str, list[dict[str, Any]], set[str], dict[str, Any]]:
    try:
        import pptx
    except ImportError as exc:  # pragma: no cover - depends on local install
        raise RuntimeError("python-pptx is required for structured PPTX conversion") from exc

    prs = pptx.Presentation(str(input_path))
    asset_dir.mkdir(parents=True, exist_ok=True)
    image_hashes: set[str] = set()
    digest_to_path: dict[str, Path] = {}
    inline_images: list[dict[str, Any]] = []
    lines: list[str] = []
    stats = {
        "slides": len(prs.slides),
        "slide_titles": 0,
        "headings": 0,
        "list_items": 0,
        "body_paragraphs": 0,
        "tables": 0,
        "charts": 0,
    }
    image_index = 0

    slide_height = int(getattr(prs.slide_height, "emu", prs.slide_height))
    for slide_number, slide in enumerate(prs.slides, start=1):
        if lines:
            lines.append("")
        lines.append(f"<!-- Slide number: {slide_number} -->")

        title_shape = choose_pptx_title_shape(slide, slide_height)
        if title_shape is not None:
            title_text = clean_text(title_shape.text)
            lines.append(f"## Slide {slide_number}: {title_text}")
            stats["slide_titles"] += 1
        else:
            lines.append(f"## Slide {slide_number}")
        lines.append("")

        for shape in iter_pptx_shapes(slide.shapes):
            if same_pptx_shape(shape, title_shape):
                continue

            if getattr(shape, "has_table", False):
                table_md = pptx_table_to_markdown(shape)
                if table_md:
                    lines.append(table_md)
                    lines.append("")
                    stats["tables"] += 1
                continue

            if getattr(shape, "has_chart", False):
                chart_title = pptx_chart_title(shape)
                lines.append(f"### Chart: {chart_title}")
                lines.append("")
                lines.append("> Chart detected. Review the original PPTX for visual chart fidelity.")
                lines.append("")
                stats["charts"] += 1
                continue

            if hasattr(shape, "image"):
                image_index += 1
                image_md, image_info = save_pptx_image(
                    shape,
                    asset_dir,
                    markdown_path,
                    image_hashes,
                    digest_to_path,
                    image_index,
                )
                image_info["slide"] = slide_number
                inline_images.append(image_info)
                lines.append(image_md)
                lines.append("")
                continue

            text_lines, text_stats = pptx_text_shape_to_markdown(
                shape,
                is_title_shape=False,
                list_mode=list_mode,
            )
            if text_lines:
                lines.extend(text_lines)
                lines.append("")
                for key in ("headings", "list_items", "body_paragraphs"):
                    stats[key] += text_stats[key]

        if slide.has_notes_slide:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                lines.append("### Speaker Notes")
                lines.append("")
                for note_line in notes.splitlines():
                    note_line = clean_text(note_line)
                    if note_line:
                        lines.append(f"> {note_line}")
                lines.append("")

    markdown = "\n".join(lines).rstrip() + "\n"
    metadata = {
        "structured": True,
        "list_mode": list_mode,
        "stats": stats,
    }
    return markdown, inline_images, image_hashes, metadata


def save_data_uri_images(
    markdown: str, asset_dir: Path, markdown_path: Path
) -> tuple[str, list[dict[str, Any]], set[str]]:
    asset_dir.mkdir(parents=True, exist_ok=True)
    images: list[dict[str, Any]] = []
    seen_hashes: set[str] = set()
    index = 0

    def replace(match: re.Match[str]) -> str:
        nonlocal index
        alt = match.group(1).strip()
        mime = match.group(2).lower()
        b64 = re.sub(r"\s+", "", match.group(3))
        raw = base64.b64decode(b64, validate=False)
        digest = sha256_bytes(raw)
        seen_hashes.add(digest)

        index += 1
        ext = MIME_EXTENSION.get(mime, ".bin")
        alt_stem = Path(alt).stem if alt else ""
        name_stem = sanitize_filename(alt_stem, f"image-{index:03d}")
        if name_stem.lower() in {"image", "picture", "media"}:
            name_stem = f"{name_stem}-{index:03d}"
        out_path = unique_path(asset_dir / f"{name_stem}{ext}")
        out_path.write_bytes(raw)
        link = rel_link(out_path, markdown_path)
        images.append(
            {
                "alt": alt,
                "mime": mime,
                "path": str(out_path),
                "markdown_link": link,
                "bytes": len(raw),
                "sha256": digest,
            }
        )
        return f"![{alt}]({link})"

    rewritten = DATA_IMAGE_RE.sub(replace, markdown)
    return rewritten, images, seen_hashes


def extract_package_images(
    input_path: Path, asset_dir: Path, markdown_path: Path, skip_hashes: set[str]
) -> list[dict[str, Any]]:
    prefix = ZIP_IMAGE_PREFIX.get(input_path.suffix.lower())
    if prefix is None or not zipfile.is_zipfile(input_path):
        return []

    extracted: list[dict[str, Any]] = []
    asset_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(input_path) as archive:
        names = sorted(archive.namelist())
        image_index = 0
        for name in names:
            lower = name.lower()
            if not lower.startswith(prefix):
                continue
            ext = Path(lower).suffix
            if ext not in IMAGE_EXTENSIONS:
                continue
            raw = archive.read(name)
            digest = sha256_bytes(raw)
            if digest in skip_hashes:
                continue
            skip_hashes.add(digest)
            image_index += 1
            source_name = sanitize_filename(Path(name).name, f"embedded-{image_index:03d}{ext}")
            if not Path(source_name).suffix:
                source_name += ext
            out_path = unique_path(asset_dir / f"embedded-{image_index:03d}-{source_name}")
            out_path.write_bytes(raw)
            extracted.append(
                {
                    "source": name,
                    "path": str(out_path),
                    "markdown_link": rel_link(out_path, markdown_path),
                    "bytes": len(raw),
                    "sha256": digest,
                    "position": "unplaced",
                }
            )
    return extracted


def inspect_ooxml_features(input_path: Path) -> list[str]:
    suffix = input_path.suffix.lower()
    if suffix not in ZIP_IMAGE_PREFIX or not zipfile.is_zipfile(input_path):
        return []

    features: set[str] = set()
    with zipfile.ZipFile(input_path) as archive:
        names = set(archive.namelist())
        lowered = {name.lower() for name in names}

        if suffix == ".docx":
            checks = {
                "headers": "word/header",
                "footers": "word/footer",
                "comments": "word/comments",
                "footnotes": "word/footnotes",
                "endnotes": "word/endnotes",
                "charts": "word/charts/",
                "smartart_or_diagrams": "word/diagrams/",
                "embedded_objects": "word/embeddings/",
                "tracked_changes_or_revision_metadata": "word/people.xml",
            }
            document_xml = archive.read("word/document.xml").decode("utf-8", "ignore") if "word/document.xml" in names else ""
            if "<m:oMath" in document_xml or "<m:oMathPara" in document_xml:
                features.add("equations")
            if "<w:ins" in document_xml or "<w:del" in document_xml:
                features.add("tracked_changes")
        elif suffix == ".pptx":
            checks = {
                "charts": "ppt/charts/",
                "smartart_or_diagrams": "ppt/diagrams/",
                "embedded_objects": "ppt/embeddings/",
            }
        else:
            checks = {
                "charts": "xl/charts/",
                "drawings_or_positioned_images": "xl/drawings/",
                "comments": "xl/comments",
                "pivot_tables": "xl/pivottables/",
                "external_links": "xl/externallinks/",
            }

        for label, needle in checks.items():
            if any(name.startswith(needle) for name in lowered):
                features.add(label)

    return sorted(features)


def append_unplaced_images(
    markdown: str, supplemental_images: list[dict[str, Any]]
) -> str:
    if not supplemental_images:
        return markdown
    lines = [
        markdown.rstrip(),
        "",
        "## Extracted Images",
        "",
        "<!-- These images were extracted from the Office package; MarkItDown did not expose exact inline positions. -->",
        "",
    ]
    for index, image in enumerate(supplemental_images, start=1):
        lines.append(f"![embedded image {index:03d}]({image['markdown_link']})")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def build_warnings(
    unsupported_features: list[str],
    supplemental_images: list[dict[str, Any]],
    remaining_data_uris: int,
) -> list[str]:
    warnings: list[str] = [
        "Markdown cannot preserve exact Office visual layout, fonts, colors, absolute positioning, animations, or all spreadsheet styling."
    ]
    if unsupported_features:
        warnings.append(
            "Potentially lossy Office features detected: " + ", ".join(unsupported_features)
        )
    if supplemental_images:
        warnings.append(
            f"{len(supplemental_images)} image(s) were extracted without a reliable inline position."
        )
    if remaining_data_uris:
        warnings.append(f"{remaining_data_uris} data:image URI(s) remain in Markdown.")
    return warnings


def main() -> int:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        print(f"error: input does not exist: {input_path}", file=sys.stderr)
        return 1
    suffix = input_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        print(
            f"error: unsupported extension {suffix}; expected {sorted(SUPPORTED_EXTENSIONS)}",
            file=sys.stderr,
        )
        return 1

    markdown_path = (
        Path(args.out).expanduser().resolve()
        if args.out
        else input_path.with_suffix(".md").resolve()
    )
    asset_dir = (
        Path(args.asset_dir).expanduser().resolve()
        if args.asset_dir
        else markdown_path.with_name(f"{markdown_path.stem}_assets").resolve()
    )
    report_path = (
        Path(args.report_out).expanduser().resolve()
        if args.report_out
        else markdown_path.with_suffix(markdown_path.suffix + ".report.json").resolve()
    )

    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    asset_dir.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)

    structural_metadata: dict[str, Any] | None = None
    if suffix == ".pptx" and args.pptx_engine == "structured":
        markdown, inline_images, image_hashes, structural_metadata = pptx_structured_convert(
            input_path,
            asset_dir,
            markdown_path,
            args.pptx_list_mode,
        )
        engine = {
            "method": "python-pptx-structured",
            "path": "pptx.Presentation",
            "list_mode": args.pptx_list_mode,
        }
    else:
        raw_markdown, engine = markitdown_convert(input_path)
        markdown, inline_images, image_hashes = save_data_uri_images(
            raw_markdown, asset_dir, markdown_path
        )
    supplemental_images = extract_package_images(
        input_path, asset_dir, markdown_path, image_hashes
    )
    if not args.no_append_unplaced:
        markdown = append_unplaced_images(markdown, supplemental_images)

    markdown = markdown.rstrip() + "\n"
    markdown_path.write_text(markdown, encoding="utf-8")

    remaining_data_uris = len(re.findall(r"data:image/[A-Za-z0-9.+-]+;base64,", markdown))
    unsupported_features = inspect_ooxml_features(input_path)
    warnings = build_warnings(unsupported_features, supplemental_images, remaining_data_uris)

    strict_failures: list[str] = []
    if unsupported_features:
        strict_failures.append("unsupported Office features detected")
    if supplemental_images:
        strict_failures.append("images extracted without reliable inline positions")
    if remaining_data_uris:
        strict_failures.append("data:image URIs remain")

    report = {
        "input": str(input_path),
        "output": str(markdown_path),
        "asset_dir": str(asset_dir),
        "report": str(report_path),
        "engine": engine,
        "counts": {
            "inline_images": len(inline_images),
            "supplemental_unplaced_images": len(supplemental_images),
            "remaining_data_uris": remaining_data_uris,
            "markdown_chars": len(markdown),
            "markdown_lines": len(markdown.splitlines()),
        },
        "inline_images": inline_images,
        "supplemental_unplaced_images": supplemental_images,
        "structure": structural_metadata,
        "unsupported_features": unsupported_features,
        "warnings": warnings,
        "strict": {
            "enabled": bool(args.strict),
            "passed": not strict_failures,
            "failures": strict_failures,
        },
    }
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    summary = {
        "ok": not (args.strict and strict_failures),
        "output": str(markdown_path),
        "asset_dir": str(asset_dir),
        "report": str(report_path),
        "inline_images": len(inline_images),
        "supplemental_unplaced_images": len(supplemental_images),
        "structure": structural_metadata,
        "unsupported_features": unsupported_features,
        "strict_failures": strict_failures if args.strict else [],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 2 if args.strict and strict_failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
